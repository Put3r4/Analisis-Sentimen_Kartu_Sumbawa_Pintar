"""Modul Preprocessing Data Komentar Kartu Sumbawa Pintar (KSP).

Modul ini bertanggung jawab untuk membaca data mentah hasil scraping dari Apify,
membersihkan komentar kosong, mengonversi format waktu ke WITA (Asia/Makassar),
menyesuaikan struktur 14 kolom terstandar, melakukan deduplikasi bertingkat,
serta meng-update dataset fix dan data labelling.
"""

import glob
from pathlib import Path
from typing import Dict

import numpy as np
import openpyxl
import pandas as pd

# Kamus pemetaan URL postingan Facebook ke Akun Sumber Postingan
URL_TO_SUMBER: Dict[str, str] = {
    "https://web.facebook.com/share/v/1CbMB5dHSV/": "Liputan Sumbawa",
    "https://web.facebook.com/share/p/17tjeersj7/": "Sumbawa Infoo",
    "https://web.facebook.com/share/p/1HzTaLLRra/": "SAMAWA LAWYER CLUB",
    "https://web.facebook.com/share/p/1DQLzLvQaa/": "Syarafuddin Jarot",
    "https://web.facebook.com/share/p/14pCZ7PF3ZF/": "Abdi-nha Wibawa",
    "https://web.facebook.com/share/p/1E4KsBA74z/": "Semaras.Sia",
    "https://web.facebook.com/share/p/1J2itfBprs/": "Agustina Sukardi",
    "https://web.facebook.com/share/p/1FLq2K2Y5N/": "Sahril Imran",
    "https://web.facebook.com/share/v/1L1TkDXWXM/": "Iwan Sofian",
    "https://web.facebook.com/share/p/1D582FxEeZ/": "Prokopim Sumbawa",
    "https://web.facebook.com/share/p/1FAweV127P/": "Fajar Rachmat",
    "https://web.facebook.com/share/p/1HUK6a9QkS/": "Faisal Salim Al Atas",
    "https://web.facebook.com/share/p/1Kvm1Mxe6W/": "Nununglusiyana",
    "https://web.facebook.com/share/p/1JKSaFDCkr/": "Boerhan Samawarea",
    "https://web.facebook.com/share/p/1dxhBJCKfb/": "Belum diketahui",
    "https://web.facebook.com/share/p/1F7s4wJHmT/": "Belum diketahui",
}


def load_raw_datasets(dataset_dir: Path) -> pd.DataFrame:
    """Membaca dan menggabungkan seluruh file CSV mentah dari folder dataset.

    Args:
        dataset_dir (Path): Jalur folder tempat file CSV mentah berada.

    Returns:
        pd.DataFrame: DataFrame gabungan dari seluruh file CSV.
    """
    pattern = str(dataset_dir / "*.csv")
    csv_files = sorted(glob.glob(pattern))
    if not csv_files:
        raise FileNotFoundError(f"Tidak ada file CSV ditemukan di: {dataset_dir}")

    dfs = [pd.read_csv(filepath, low_memory=False) for filepath in csv_files]
    df_raw = pd.concat(dfs, ignore_index=True)
    print(f"Total file CSV dibaca: {len(csv_files)}")
    print(f"Total baris mentah gabungan: {len(df_raw)}")
    return df_raw


def process_dataset(df_raw: pd.DataFrame) -> pd.DataFrame:
    """Melakukan pembersihan awal, deduplikasi ID, dan transformasi format kolom.

    Args:
        df_raw (pd.DataFrame): DataFrame mentah hasil penggabungan CSV.

    Returns:
        pd.DataFrame: DataFrame terstruktur sesuai 14 kolom template KSP.
    """
    # 1. Hapus komentar tanpa teks (stiker/reaksi kosong)
    df = df_raw.dropna(subset=["text"]).copy()
    df = df[df["text"].astype(str).str.strip() != ""].copy()

    # 2. Hapus duplikasi berdasarkan ID komentar unik
    df = df.drop_duplicates(subset=["id"]).reset_index(drop=True)

    # 3. Penentuan Nama Akun Komentator
    author_name = df["author/name"] if "author/name" in df.columns else np.nan
    profile_name = df["profileName"] if "profileName" in df.columns else np.nan
    df["Nama Akun"] = author_name.fillna(profile_name).fillna("Tanpa Nama")

    # 4. Standarisasi Waktu (UTC -> WITA / Asia/Makassar)
    dt_utc = pd.to_datetime(df["date"], utc=True, errors="coerce")
    dt_wita = dt_utc.dt.tz_convert("Asia/Makassar")
    df["Tanggal"] = dt_wita.dt.strftime("%Y-%m-%d")
    df["Waktu_Komentar_WITA"] = dt_wita.dt.strftime("%Y-%m-%d %H:%M:%S")

    # 5. Normalisasi Link Sumber & Pemetaan Akun Sumber Postingan dari URL
    urls_normalized = (
        df["facebookUrl"]
        .astype(str)
        .str.replace("https://www.facebook.com/", "https://web.facebook.com/", regex=False)
    )
    df["Akun_Sumber_Postingan"] = (
        urls_normalized.map(URL_TO_SUMBER).fillna("Belum diketahui")
    )

    # 6. Deduplikasi Lanjutan: Nama Akun + Isi Komentar Clean + Tanggal (Waktu yang Sama)
    df["isi_clean"] = (
        df["text"].astype(str).str.strip().str.replace(r"\s+", " ", regex=True)
    )
    subset_dedup = ["Nama Akun", "isi_clean", "Tanggal"]
    df_dedup = df.drop_duplicates(subset=subset_dedup, keep="first").reset_index(drop=True)

    # 7. Penjenisan Komentar Utama vs Balasan & Metrik Interaksi pada df_dedup
    depth = (
        df_dedup["threadingDepth"].fillna(0)
        if "threadingDepth" in df_dedup.columns
        else 0
    )
    jenis = np.where(depth == 0, "Komentar Utama", "Balasan")

    likes = (
        df_dedup["likesCount"].fillna(0).astype(int)
        if "likesCount" in df_dedup.columns
        else 0
    )
    comments = (
        df_dedup["commentsCount"].fillna(0).astype(int)
        if "commentsCount" in df_dedup.columns
        else 0
    )

    # 8. Penyusunan DataFrame Final 14 Kolom Terstruktur
    df_final = pd.DataFrame(
        {
            "No": np.arange(1, len(df_dedup) + 1),
            "Tanggal": df_dedup["Tanggal"],
            "Platform/Sumber": "Facebook",
            "Nama Akun": df_dedup["Nama Akun"],
            "Isi Komentar": df_dedup["text"],
            "Link Sumber": df_dedup["facebookUrl"],
            "Jenis": jenis,
            "Jumlah_Like": likes,
            "Jumlah_Balasan": comments,
            "Waktu_Komentar_WITA": df_dedup["Waktu_Komentar_WITA"],
            "Akun_Sumber_Postingan": df_dedup["Akun_Sumber_Postingan"],
            "Kategori Isu": "",
            "Sentimen": "",
            "Catatan Analis": "",
        }
    )

    df_final = df_final.sort_values(
        ["Akun_Sumber_Postingan", "Waktu_Komentar_WITA"]
    ).reset_index(drop=True)
    df_final["No"] = np.arange(1, len(df_final) + 1)
    print(f"Total baris bersih setelah deduplikasi: {len(df_final)}")
    return df_final


def deduplicate_labelling_dataset(excel_path: Path) -> pd.DataFrame:
    """Melakukan deduplikasi komentar pada dataset labelling.

    Deduplikasi dilakukan untuk komentar dengan kombinasi Nama Akun, Isi Komentar,
    dan Tanggal yang sama. Baris yang sudah memiliki label (Kategori Isu/Sentimen)
    diprioritaskan untuk dipertahankan.

    Args:
        excel_path (Path): Jalur file Excel Data_Komentar_KSP_labelling.xlsx.

    Returns:
        pd.DataFrame: DataFrame labelling yang sudah dideduplikasi.
    """
    if not excel_path.exists():
        raise FileNotFoundError(f"File labelling tidak ditemukan: {excel_path}")

    df_lab = pd.read_excel(excel_path, sheet_name="Data Komentar")
    initial_count = len(df_lab)

    df_lab["isi_clean"] = (
        df_lab["Isi Komentar"].astype(str).str.strip().str.replace(r"\s+", " ", regex=True)
    )
    df_lab["has_label"] = (
        df_lab["Kategori Isu"].notna() | df_lab["Sentimen"].notna()
    )

    # Prioritaskan baris berlabel, lalu urutan awal
    df_sorted = df_lab.sort_values(by=["has_label", "No"], ascending=[False, True])
    subset = ["Nama Akun", "isi_clean", "Tanggal"]

    df_dedup = df_sorted.drop_duplicates(subset=subset, keep="first").copy()
    df_dedup = df_dedup.sort_values(
        by=["Akun_Sumber_Postingan", "Waktu_Komentar_WITA"]
    ).reset_index(drop=True)

    df_dedup["No"] = np.arange(1, len(df_dedup) + 1)
    df_dedup = df_dedup.drop(columns=["isi_clean", "has_label"])

    final_count = len(df_dedup)
    print(f"Labelling Data Initial: {initial_count} baris")
    print(
        f"Labelling Data Final: {final_count} baris (Dihapus {initial_count - final_count} duplikat)"
    )

    # Simpan kembali ke Excel dengan meng-update sheet 'Data Komentar' saja
    wb = openpyxl.load_workbook(excel_path)
    if "Data Komentar" in wb.sheetnames:
        del wb["Data Komentar"]
    ws = wb.create_sheet("Data Komentar", index=1)

    ws.append(list(df_dedup.columns))
    for row in df_dedup.itertuples(index=False):
        ws.append(list(row))

    wb.save(excel_path)
    return df_dedup


def save_processed_dataset(df_final: pd.DataFrame, output_dir: Path) -> None:
    """Menyimpan dataset hasil olahan ke format Excel dan CSV.

    Args:
        df_final (pd.DataFrame): DataFrame final yang siap disimpan.
        output_dir (Path): Folder tujuan penyimpanan file fix.
    """
    output_dir.mkdir(parents=True, exist_ok=True)

    excel_file = output_dir / "Data_Komentar_KSP_fix.xlsx"
    csv_file = output_dir / "Data_Komentar_KSP_fix.csv"

    with pd.ExcelWriter(excel_file, engine="openpyxl") as writer:
        df_final.to_excel(writer, sheet_name="Data Komentar", index=False)

    df_final.to_csv(csv_file, index=False, encoding="utf-8-sig")
    print(f"Dataset fix disimpan ke:\n - {excel_file}\n - {csv_file}")


def main() -> None:
    """Fungsi utama eksekusi pipeline preprocessing data."""
    base_dir = Path(__file__).resolve().parent.parent
    raw_dir = base_dir / "data_scraping" / "dataset"
    out_dir = base_dir / "data_scraping" / "fix_dataset"
    labelling_file = (
        base_dir / "data_scraping" / "data_labelling" / "Data_Komentar_KSP_labelling.xlsx"
    )

    print("=== TAHAP 1: PREPROCESSING DATASET MENTAH ===")
    df_raw = load_raw_datasets(raw_dir)
    df_processed = process_dataset(df_raw)
    save_processed_dataset(df_processed, out_dir)

    print("\n=== TAHAP 2: DEDUPLIKASI DATASET LABELLING ===")
    deduplicate_labelling_dataset(labelling_file)


if __name__ == "__main__":
    main()
